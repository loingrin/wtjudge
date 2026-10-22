from __future__ import annotations

import csv
import html
import base64
import sys
from collections import defaultdict
from dataclasses import asdict, fields
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from wtjudge.core.models import PointResult, STATUS_TEXT
from wtjudge.presentation import STATUS_COLOURS, score_colours, uniform_error
from wtjudge.storage.project import atomic_json
from .style import REPORT_CSS

NOTICE = "Windows 本地模拟评测结果；墙钟时间及 MinGW 环境与 CSP 官方 Linux 评测可能有差异。"
WARNING = "结果不完整 / 不可视作最终成绩；未完成项不作为正式 0 分，不提供正式排名。"


def _submission_display(batch, row, points, problem):
    """Keep unfinished submissions distinct from homogeneous final errors."""
    identities = [(point["student_id"], point["problem_id"], point["testcase_id"]) for point in points]
    complete = bool(points) and len(identities) == len(set(identities)) and all(
        point["status"] in STATUS_TEXT and point["status"] not in {"SYSTEM_ERROR", "CANCELLED", "AMBIGUOUS_SOURCE"}
        for point in points)
    if batch.testcase_ids:
        expected = batch.testcase_ids.get(problem, [])
        complete = complete and bool(expected) and len(expected) == len(set(expected)) and {
            identity[2] for identity in identities} == set(expected)
    else:
        expected = None
        # A legacy report has no per-problem manifest. Its completed row is the
        # only evidence that all expected points have arrived.
        complete = complete and row["completed"]
    return (uniform_error(points, expected) if complete else ""), complete


def _excel_colours(cell, colours, *, bold=False):
    background, foreground = colours
    cell.fill = PatternFill("solid", fgColor=background.lstrip("#"))
    cell.font = Font(name="Microsoft YaHei UI", size=10, color=foreground.lstrip("#"), bold=bold)


def coverage_text(batch):
    invalid = sum(p.status in {"SYSTEM_ERROR", "CANCELLED", "AMBIGUOUS_SOURCE"} for p in batch.points)
    valid = sum(p.status not in {"SYSTEM_ERROR", "CANCELLED", "AMBIGUOUS_SOURCE"} for p in batch.points)
    missing = max(0, batch.expected_points - len(batch.points))
    return f"预计 {batch.expected_points} 项；已判定 {valid} 项；无效/取消 {invalid} 项；缺少记录 {missing} 项。"


def append_excel(sheet, values):
    """Sanitize before Cell construction; openpyxl rejects NUL on append."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    values = [ILLEGAL_CHARACTERS_RE.sub("", value).replace("\ufffe", "").replace("\uffff", "")[:32000]
              if isinstance(value, str) else value for value in values]
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = "s"


def export_reports(batch, output_root):
    """Always create a fresh directory; previous reports are never overwritten."""
    root = Path(output_root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    directory = root / f"wtjudge-{stamp}"
    directory.mkdir(exist_ok=False)
    count_fields = [f"{s}_count" for s in STATUS_TEXT if s != "AC"]
    used = {"rank", "student_id", "total", "completed", "ac_count", *count_fields}
    columns = {}
    for problem in batch.problems:
        column = problem
        while column in used:
            column = "problem:" + column
        used.add(column)
        columns[problem] = column
    score_fields = ["rank", "student_id", *columns.values(), "total", "completed", "ac_count", *count_fields]
    rows = []
    standings = batch.standings()
    for row in standings:
        record = {key: row[key] for key in ("rank", "student_id", "total", "completed", "ac_count")}
        record.update({columns[p]: row["scores"][p] for p in batch.problems})
        record.update({f"{s}_count": n for s, n in row["error_counts"].items()})
        rows.append(record)
    detail_fields = [f.name for f in fields(PointResult)]
    details = [asdict(p) for p in batch.points]
    by_submission = defaultdict(list)
    for point in details:
        by_submission[point["student_id"], point["problem_id"]].append(point)
    total_full_score = sum(batch.full_scores.values())
    for name, headers, records in (("scores.csv", score_fields, rows), ("details.csv", detail_fields, details)):
        with open(directory / name, "w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.DictWriter(stream, fieldnames=headers)
            writer.writeheader()
            writer.writerows(records)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, headers, records in (("总成绩", score_fields, rows), ("测试点详情", detail_fields, details)):
        sheet = workbook.create_sheet(name)
        append_excel(sheet, headers)
        for record in records:
            append_excel(sheet, [record.get(key) for key in headers])
        sheet.freeze_panes = "C2" if name == "总成绩" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 95
        sheet.sheet_properties.tabColor = "9CB68A"
        sheet.row_dimensions[1].height = 30
        for cell in sheet[1]:
            cell.font = Font(name="Microsoft YaHei UI", size=10, bold=True, color="52664A")
            cell.fill = PatternFill("solid", fgColor="EEF3E9")
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="thin", color="DEE7D7"))
        for column, header in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column)].width = 65 if header == "message" else max(14, min(28, len(header) + 4))
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height = 26
            for cell in row:
                # Student-controlled text must remain text, not an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                    cell.value = "".join(c for c in cell.value if ord(c) >= 32 or c in "\t\n\r")[:32000]
                cell.font = Font(name="Microsoft YaHei UI", size=10, color="4B5547")
                cell.alignment = Alignment(vertical="center")
                cell.fill = PatternFill("solid", fgColor="FFFFFF" if cell.row % 2 == 0 else "F8FAF6")
                cell.border = Border(bottom=Side(style="hair", color="EEF1EA"))
            if name == "总成绩":
                standing = standings[row[0].row - 2]
                for problem in batch.problems:
                    cell = row[headers.index(columns[problem])]
                    error, complete = _submission_display(batch, standing, by_submission[standing["student_id"], problem], problem)
                    colours = STATUS_COLOURS[error] if error else score_colours(cell.value, batch.full_scores[problem])
                    _excel_colours(cell, colours)
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if error:
                        # Excel renders the verdict while keeping the score a
                        # numeric value for SUM, sorting and external readers.
                        cell.number_format = f'"{"缺源码" if error == "NO_SOURCE" else error}"'
                        cell.comment = Comment(f"{error} · {STATUS_TEXT[error]}\n全部测试点均为 {error}；原始得分：{cell.value}", "wtjudge")
                    elif not complete:
                        cell.comment = Comment("未完成 / 无效项；此数值不能视作正式成绩。", "wtjudge")
                total_cell = row[headers.index("total")]
                _excel_colours(total_cell, score_colours(total_cell.value, total_full_score), bold=True)
                total_cell.alignment = Alignment(horizontal="center", vertical="center")
                if not row[headers.index("completed")].value:
                    row[headers.index("completed")].font = Font(name="Microsoft YaHei UI", size=10, color="AC8758")
            if name == "测试点详情":
                status_cell = row[headers.index("status")]
                _excel_colours(status_cell, STATUS_COLOURS.get(status_cell.value, ("#fff2df", "#ae8753")), bold=True)
                status_cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    meta = workbook.create_sheet("说明")
    append_excel(meta, [batch.project_name])
    append_excel(meta, ["完整成绩" if batch.complete else WARNING])
    append_excel(meta, [coverage_text(batch)])
    if batch.failure:
        append_excel(meta, ["批次错误", batch.failure])
    if hasattr(batch, "coverage_errors"):
        for error in batch.coverage_errors():
            append_excel(meta, ["覆盖检查", error])
    for warning in getattr(batch, "cleanup_warnings", []):
        append_excel(meta, ["临时文件清理警告", warning])
    append_excel(meta, [NOTICE])
    append_excel(meta, [batch.started_at, batch.finished_at])
    meta.column_dimensions["A"].width = 100
    meta.column_dimensions["B"].width = 85
    meta.sheet_view.showGridLines = False
    meta.sheet_properties.tabColor = "DCE5D2"
    for row in meta:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei UI", size=10, color="849079")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        meta.row_dimensions[row[0].row].height = 28
    meta["A1"].font = Font(name="Microsoft YaHei UI", size=15, color="4E6C37", bold=True)
    workbook.save(directory / "scores.xlsx")
    atomic_json(directory / "results.json", batch.to_dict())
    _html_report(batch, standings, details, directory / "report.html")
    return str(directory)


def _logo_data_url():
    roots = [Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[2])) / "assets",
             Path(__file__).resolve().parents[2] / "assets"]
    for root in roots:
        logo = root / "wotu.png"
        if logo.is_file():
            return "data:image/png;base64," + base64.b64encode(logo.read_bytes()).decode("ascii")
    return ""


def _html_report(batch, rows, details, path):
    esc = lambda value: html.escape(str(value), quote=True)
    grouped = defaultdict(list)
    by_submission = defaultdict(list)
    for point in details:
        grouped[point["student_id"]].append(point)
        by_submission[point["student_id"], point["problem_id"]].append(point)
    row_by_student = {row["student_id"]: row for row in rows}
    anchors = {student: f"student-{i}" for i, student in enumerate(batch.students)}
    total_full_score = sum(batch.full_scores.values())
    with open(path, "w", encoding="utf-8") as stream:
        stream.write('<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">')
        stream.write(f'<title>{esc(batch.project_name)} · 成绩</title><style>{REPORT_CSS}</style></head><body><div class="shell"><aside class="sidebar"><a class="brand" href="#scores">')
        logo = _logo_data_url()
        stream.write(f'<img src="{logo}" alt="" width="32" height="32">' if logo else '<span class="brand-mark">w</span>')
        stream.write('<span>wtjudge</span></a><nav aria-label="报告导航"><a href="#scores"><span class="nav-number">01</span>总成绩</a><a href="#points"><span class="nav-number">02</span>测试点详情</a><a href="#method"><span class="nav-number">03</span>评测说明</a></nav></aside><main>')
        completed_text = "已完成" if batch.complete else "结果不完整"
        status_class = "completion" if batch.complete else "completion partial"
        date = batch.finished_at or batch.started_at
        try:
            date = datetime.fromisoformat(date).strftime("%Y.%m.%d %H:%M")
        except (TypeError, ValueError):
            pass
        stream.write(f'<header class="topline"><div><h1>{esc(batch.project_name)}</h1><p class="subline">{len(batch.students)} 名考生 · {len(batch.problems)} 道题 · {esc(date)}</p></div><span class="{status_class}">{completed_text}</span></header>')
        if not batch.complete:
            stream.write(f'<div class="warning"><strong>{esc(WARNING)}</strong>{esc(batch.failure)}</div>')
        for error in batch.coverage_errors() if hasattr(batch, "coverage_errors") else []:
            stream.write(f'<p class="warning">{esc(error)}</p>')
        for warning in getattr(batch, "cleanup_warnings", []):
            stream.write(f'<p class="warning">临时文件清理警告：{esc(warning)}</p>')
        stream.write('<section id="scores" class="surface"><div class="section-bar"><h2>总成绩</h2><label class="visually-hidden" for="search">准考证号筛选</label><input class="search" id="search" placeholder="搜索考生" autocomplete="off"></div><div class="table-wrap"><table><thead><tr><th class="rank">名次</th><th>准考证号</th>')
        for problem in batch.problems:
            stream.write(f'<th class="numeric">{esc(problem)}</th>')
        stream.write('<th class="numeric">总分</th><th>状态</th></tr></thead><tbody>')
        for row in rows:
            student = row["student_id"]
            stream.write(f'<tr data-student="{esc(student)}"><td class="rank">{esc(row["rank"])}</td><td class="student"><a class="student-link" href="#{anchors[student]}">{esc(student)}</a></td>')
            for problem in batch.problems:
                points = by_submission[student, problem]
                score = row["scores"][problem]
                error, complete = _submission_display(batch, row, points, problem)
                background, foreground = STATUS_COLOURS[error] if error else score_colours(score, batch.full_scores[problem])
                display = "缺源码" if error == "NO_SOURCE" else error or score
                css = "uniform-error" if error else "" if complete else "invalid"
                note = "" if complete else "未完成"
                title = f"{STATUS_TEXT[error]}；得分 {score} / {batch.full_scores[problem]}" if error else f"{score} / {batch.full_scores[problem]}"
                note_html = f'<span class="cell-note">{note}</span>' if note else ""
                stream.write(f'<td class="numeric score {css}" data-score="{score}" title="{esc(title)}" style="background:{background};color:{foreground}">{esc(display)}{note_html}</td>')
            state = "已判定" if row["completed"] else "未完成 / 无效项"
            background, foreground = score_colours(row["total"], total_full_score)
            stream.write(f'<td class="numeric total" style="background:{background};color:{foreground}">{row["total"]}</td><td class="row-state {"" if row["completed"] else "incomplete"}">{state}</td></tr>')
        if not rows:
            stream.write(f'<tr><td class="empty" colspan="{len(batch.problems) + 4}">没有成绩记录</td></tr>')
        stream.write(f'</tbody></table></div><p class="coverage">{esc(coverage_text(batch))}</p></section><div class="details-heading" id="points"><h2>测试点详情</h2><span>{len(details)} 项</span></div>')
        for student in batch.students:
            row = row_by_student[student]
            points = grouped[student]
            count_ac = sum(point["status"] == "AC" for point in points)
            stream.write(f'<details class="student-detail" id="{anchors[student]}" data-student="{esc(student)}"><summary><span>{esc(student)}</span><span class="detail-stats">{row["total"]} 分 · AC {count_ac}/{len(points)}</span></summary><div class="detail-inner table-wrap"><table class="point-table"><thead><tr><th>题目</th><th>测试点</th><th>状态</th><th class="numeric">得分</th><th class="numeric">墙钟 ms</th><th class="numeric">峰值 KB</th><th>诊断</th></tr></thead><tbody>')
            for point in points:
                status = point["status"]
                css = status.lower() if status in STATUS_TEXT else "unknown"
                message = f'退出码：{point["exit_code"]}\n{point["message"]}'
                stream.write(f'<tr><td>{esc(point["problem_id"])}</td><td>{esc(point["testcase_id"])}</td><td><span class="verdict {css}" title="{esc(STATUS_TEXT.get(status, status))}">{esc(status)}</span></td><td class="numeric">{point["score"]}<span class="subline"> / {point["max_score"]}</span></td><td class="numeric">{round(point["time_ms"], 2)}</td><td class="numeric">{point["peak_memory_kb"]}</td><td class="message"><details class="diagnostic"><summary>{esc(STATUS_TEXT.get(status, status))}</summary><pre>{esc(message)}</pre></details></td></tr>')
            stream.write('</tbody></table></div></details>')
        stream.write(f'<details class="method" id="method"><summary>评测说明</summary><p>{esc(NOTICE)}</p><p>AC 获得该点满分；缺源码、编译错误无需执行程序。系统错误及取消项不作为正式 0 分，含无效项的批次不提供正式排名。</p></details></main></div>')
        stream.write('''<script>document.getElementById('search').addEventListener('input',function(){const q=this.value.toLocaleLowerCase();document.querySelectorAll('[data-student]').forEach(e=>e.hidden=!e.dataset.student.toLocaleLowerCase().includes(q))});document.querySelectorAll('.student-link').forEach(a=>a.addEventListener('click',()=>{const target=document.getElementById(a.hash.slice(1));if(target)target.open=true}));</script></body></html>''')
