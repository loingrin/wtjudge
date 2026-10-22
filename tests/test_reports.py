import csv
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook

from wtjudge.core.models import BatchResult, PointResult
from wtjudge.reports.exporter import export_reports
from wtjudge.storage.project import load_result


class ResourceParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.resources = []

    def handle_starttag(self, tag, attrs):
        for key, value in attrs:
            if key == "src" or (key == "href" and tag == "link"):
                self.resources.append(value)


class ScoreParser(HTMLParser):
    """Read rendered score cells without relying on whitespace or attr order."""
    def __init__(self):
        super().__init__()
        self.scores = []
        self.totals = []
        self.current = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        classes = attrs.get("class", "").split()
        if tag == "td" and ("score" in classes or "total" in classes):
            self.current = {"attrs": attrs, "text": ""}
            (self.scores if "score" in classes else self.totals).append(self.current)

    def handle_data(self, data):
        if self.current is not None:
            self.current["text"] += data

    def handle_endtag(self, tag):
        if tag == "td":
            self.current = None


def test_reports_match_unicode_dynamic_and_offline(tmp_path):
    batch = BatchResult("中文<考试>", ["0001", "0002"], ["一", "二"], {"一": 100, "二": 100}, "2026-10-21", 4,
                        points=[PointResult(s, p, "test_1", "AC", 100, 100) for s in ["0001", "0002"] for p in ["一", "二"]],
                        finished_at="2026-10-21")
    directory = __import__("pathlib").Path(export_reports(batch, tmp_path))
    assert (directory / "scores.csv").read_bytes().startswith(b"\xef\xbb\xbf")
    with (directory / "scores.csv").open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows[0]["student_id"] == "0001"
    assert rows[0]["total"] == "200" and rows[0]["rank"] == rows[1]["rank"] == "1"
    book = load_workbook(directory / "scores.xlsx")
    assert book["总成绩"].cell(2, 2).value == "0001"
    assert book["总成绩"].freeze_panes == "C2"
    assert len(list(book["测试点详情"].rows)) == 5
    text = (directory / "report.html").read_text(encoding="utf-8")
    assert "中文&lt;考试&gt;" in text
    parser = ResourceParser()
    parser.feed(text)
    assert all(resource.startswith("data:") for resource in parser.resources)
    assert "https://" not in text and "@import" not in text
    assert load_result(directory / "results.json").complete
    batch.cancelled = True
    batch.points[-1].status = "CANCELLED"
    partial = __import__("pathlib").Path(export_reports(batch, tmp_path))
    assert partial != directory
    assert "不可视作最终成绩" in (partial / "report.html").read_text(encoding="utf-8")


def test_problem_names_cannot_overwrite_total_or_student_id(tmp_path):
    batch = BatchResult("names", ["001"], ["total", "student_id"], {"total":100, "student_id":100}, "today", 2,
                        points=[PointResult("001","total","1","AC",100,100), PointResult("001","student_id","1","WA",0,100)], finished_at="today")
    row = batch.standings()[0]
    assert row["student_id"] == "001" and row["scores"] == {"total":100, "student_id":0}
    directory = __import__("pathlib").Path(export_reports(batch, tmp_path))
    with (directory / "scores.csv").open(encoding="utf-8-sig", newline="") as stream:
        data = next(csv.DictReader(stream))
    assert data["student_id"] == "001" and data["problem:student_id"] == "0"
    assert data["total"] == "100" and data["problem:total"] == "100"


def test_report_exposes_failure_and_unjudged_counts(tmp_path):
    batch = BatchResult("interrupted", ["001"], ["p"], {"p":100}, "today", 2,
                        points=[PointResult("001","p","1","SYSTEM_ERROR",0,50)], finished_at="today", failure="临时程序清理被拒绝")
    batch.cleanup_warnings = ["待清理的 wtjudge 专用临时目录"]
    directory = __import__("pathlib").Path(export_reports(batch, tmp_path))
    report = (directory / "report.html").read_text(encoding="utf-8")
    assert "临时程序清理被拒绝" in report and "已判定 0 项" in report and "缺少记录 1 项" in report
    assert "待清理的 wtjudge 专用临时目录" in report
    book = load_workbook(directory / "scores.xlsx")
    values = " ".join(str(c.value) for row in book["说明"] for c in row)
    assert "临时程序清理被拒绝" in values and "无效/取消 1 项" in values


def test_binary_console_diagnostic_cannot_break_excel_export(tmp_path):
    batch = BatchResult("=local_title", ["001"], ["p"], {"p":100}, "today", 1,
                        points=[PointResult("001","p","1","OUTPUT_LIMIT",0,100,message="binary\x00\x01 output\ufffe")], finished_at="today")
    directory = __import__("pathlib").Path(export_reports(batch, tmp_path))
    book = load_workbook(directory / "scores.xlsx")
    assert book["说明"]["A1"].data_type == "s"
    assert book["测试点详情"].cell(2,10).value == "binary output"


def test_collapsed_report_preserves_diagnostic_and_escapes_markup(tmp_path):
    message = '<script>alert("student")</script>\nfirst mismatch at byte 12'
    batch = BatchResult("<考试>", ["0001"], ["p"], {"p":100}, "today", 1,
                        points=[PointResult("0001","p","1","WA",0,100,exit_code=0,message=message)], finished_at="today")
    directory = __import__("pathlib").Path(export_reports(batch, tmp_path))
    text = (directory / "report.html").read_text(encoding="utf-8")
    assert '<script>alert(' not in text
    assert '&lt;script&gt;alert(&quot;student&quot;)&lt;/script&gt;' in text
    assert 'first mismatch at byte 12' in text and '退出码：0' in text
    assert 'data-student="0001"' in text and 'student-link' in text


def test_score_backgrounds_verdicts_and_raw_exports_stay_consistent(tmp_path):
    bands = [(0, "#ffffff"), (10, "#cbeecb"), (30, "#cbeecb"),
             (40, "#9fdf9f"), (60, "#9fdf9f"), (70, "#7ed57e"),
             (90, "#7ed57e"), (100, "#5cc95c")]
    errors = {"RE": "#8e44ad", "CE": "#ff9900", "TLE": "#052242", "MLE": "#052242",
              "NO_OUTPUT": "#eee2f4", "OUTPUT_LIMIT": "#052242", "NO_SOURCE": "#f1f3f0"}
    problems = [f"score{score}" for score, _ in bands] + ["custom"] + list(errors)
    full_scores = {problem: 100 for problem in problems}
    full_scores["custom"] = 50
    points = []
    for score, _ in bands:
        problem = f"score{score}"
        points += [PointResult("001", problem, "1", "AC" if score else "WA", score, score),
                   PointResult("001", problem, "2", "WA" if score < 100 else "AC", 0, 100 - score)]
    points += [PointResult("001", "custom", "1", "AC", 15, 15), PointResult("001", "custom", "2", "WA", 0, 35)]
    for status in errors:
        points += [PointResult("001", status, str(i), status, 0, 50) for i in (1, 2)]
    batch = BatchResult("colours", ["001"], problems, full_scores, "today", len(points),
                        points=points, finished_at="today", testcase_ids={problem: ["1", "2"] for problem in problems})
    original = batch.to_dict()
    directory = Path(export_reports(batch, tmp_path))
    text = (directory / "report.html").read_text(encoding="utf-8")
    parser = ScoreParser()
    parser.feed(text)
    cells = dict(zip(problems, parser.scores))
    for score, background in bands:
        assert cells[f"score{score}"]["text"] == str(score)
        assert f"background:{background}" in cells[f"score{score}"]["attrs"]["style"]
    assert "background:#cbeecb" in cells["custom"]["attrs"]["style"]
    for status, background in errors.items():
        assert cells[status]["text"] == ("缺源码" if status == "NO_SOURCE" else status)
        assert f"background:{background}" in cells[status]["attrs"]["style"]
        assert cells[status]["attrs"]["data-score"] == "0"
        assert f".verdict.{status.lower()}{{background:{background};" in text
    assert ".verdict.ac{background:#5cc95c;color:#16221b}" in text
    assert "background:#cbeecb" in parser.totals[0]["attrs"]["style"]
    book = load_workbook(directory / "scores.xlsx")
    sheet = book["总成绩"]
    columns = {cell.value: cell.column for cell in sheet[1]}
    for score, background in bands:
        cell = sheet.cell(2, columns[f"score{score}"])
        assert cell.value == score and cell.data_type == "n"
        assert cell.fill.fgColor.rgb[-6:] == background[1:]
    assert sheet.cell(2, columns["custom"]).fill.fgColor.rgb[-6:] == "cbeecb"
    for status, background in errors.items():
        cell = sheet.cell(2, columns[status])
        assert cell.value == 0 and cell.data_type == "n"
        assert cell.fill.fgColor.rgb[-6:] == background[1:]
        assert cell.number_format == f'"{"缺源码" if status == "NO_SOURCE" else status}"'
        assert status in cell.comment.text and "原始得分：0" in cell.comment.text
    detail = book["测试点详情"]
    headers = [cell.value for cell in detail[1]]
    status_column = headers.index("status") + 1
    for row in detail.iter_rows(min_row=2):
        cell = row[status_column - 1]
        if cell.value in errors:
            assert cell.fill.fgColor.rgb[-6:] == errors[cell.value][1:]
    with (directory / "scores.csv").open(encoding="utf-8-sig", newline="") as stream:
        record = next(csv.DictReader(stream))
    assert all(record[status] == "0" for status in errors)
    assert record["score100"] == "100" and record["custom"] == "15"
    assert load_result(directory / "results.json").to_dict() == original == batch.to_dict()


def test_partial_and_mixed_errors_do_not_masquerade_as_final_verdict(tmp_path):
    problems = ["partial", "duplicate", "foreign", "mixed"]
    points = [PointResult("001", "partial", "1", "RE", 0, 50),
              PointResult("001", "duplicate", "1", "RE", 0, 50),
              PointResult("001", "duplicate", "1", "RE", 0, 50),
              PointResult("001", "foreign", "1", "RE", 0, 50),
              PointResult("001", "foreign", "3", "RE", 0, 50),
              PointResult("001", "mixed", "1", "RE", 0, 50),
              PointResult("001", "mixed", "2", "WA", 0, 50)]
    batch = BatchResult("partial", ["001"], problems, dict.fromkeys(problems, 100), "today", 8,
                        points=points, finished_at="today", testcase_ids=dict.fromkeys(problems, ["1", "2"]))
    directory = Path(export_reports(batch, tmp_path))
    text = (directory / "report.html").read_text(encoding="utf-8")
    parser = ScoreParser()
    parser.feed(text)
    for cell in parser.scores[:3]:
        assert cell["text"] == "0未完成"
        assert "uniform-error" not in cell["attrs"]["class"]
        assert "background:#ffffff" in cell["attrs"]["style"]
    assert parser.scores[3]["text"] == "0"
    assert "不可视作最终成绩" in text
    book = load_workbook(directory / "scores.xlsx")
    assert "未完成" in book["总成绩"].cell(2, 3).comment.text


def test_legacy_verdict_requires_completed_row(tmp_path):
    batch = BatchResult("legacy", ["001"], ["p"], {"p": 100}, "today", 2,
                        points=[PointResult("001", "p", "1", "RE", 0, 50)], finished_at="today")
    partial = Path(export_reports(batch, tmp_path))
    parser = ScoreParser()
    parser.feed((partial / "report.html").read_text(encoding="utf-8"))
    assert parser.scores[0]["text"] == "0未完成"
    batch.points.append(PointResult("001", "p", "2", "RE", 0, 50))
    complete = Path(export_reports(batch, tmp_path))
    parser = ScoreParser()
    parser.feed((complete / "report.html").read_text(encoding="utf-8"))
    assert parser.scores[0]["text"] == "RE"
